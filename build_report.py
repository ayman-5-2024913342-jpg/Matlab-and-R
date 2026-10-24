#!/usr/bin/env python3
"""
Build a Word report from a folder of MATLAB assignments.

Layout per assignment:
    Assignment N
        Code    (each .m file, monospace)
        Output  (text output + figures)

Usage examples
--------------
# From GitHub (clones into ./_repo_tmp):
python build_report.py --url https://github.com/ayman-5-2024913342-jpg/Matlab-and-R

# From a local clone:
python build_report.py --repo "C:/Users/me/Matlab-and-R"

# Also RUN the code to capture output (needs MATLAB or Octave on PATH):
python build_report.py --repo ./Matlab-and-R --run matlab
python build_report.py --repo ./Matlab-and-R --run octave

Setup:  pip install python-docx pygments
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from pygments.lexers import MatlabLexer
from pygments.token import Token
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

IMG_EXT = {".png", ".jpg", ".jpeg"}
OUT_TXT_HINTS = ("output", "result", "log")


def natural_key(p):
    s = p.name if isinstance(p, Path) else str(p)
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


# ---------------------------------------------------------------- docx helpers
def shade(paragraph, fill="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


ANSI_RE = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text):
    """Remove characters Word/XML can't store (control chars, ANSI colours, backspaces)."""
    text = ANSI_RE.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # apply backspaces (e.g. progress-bar style output)
    while "\x08" in text:
        text = re.sub(r"[^\x08]\x08", "", text, count=1) if re.search(r"[^\x08]\x08", text) else text.replace("\x08", "")
    return CTRL_RE.sub("", text)


def add_mono_block(doc, text, size=9, fill="F2F2F2"):
    """Add text as one paragraph per line, monospace, light-grey background."""
    text = clean_text(text)
    lines = text.replace("\t", "    ").splitlines() or [""]
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(size)
        shade(p, fill)


# MATLAB-editor-style colours
def token_color(tt):
    if tt in Token.Comment:
        return "008000"      # green
    if tt in Token.Keyword:
        return "0000FF"      # blue
    if tt in Token.String:
        return "A020F0"      # purple
    if tt in Token.Name.Builtin:
        return "1F5FBF"      # built-in functions
    if tt in Token.Number:
        return "B5651D"
    if tt in Token.Error:
        return None
    return None


def add_highlighted_code(doc, text, size=9, fill="F7F7F7"):
    """MATLAB code, one paragraph per line, syntax-coloured."""
    text = clean_text(text).replace("\t", "    ")
    lines, cur = [], []
    for tt, val in MatlabLexer().get_tokens(text):
        color = token_color(tt)
        italic = tt in Token.Comment
        parts = val.split("\n")
        for j, part in enumerate(parts):
            if j > 0:
                lines.append(cur)
                cur = []
            if part:
                cur.append((part, color, italic))
    lines.append(cur)
    if lines and lines[-1] == [] and len(lines) > 1:
        lines.pop()  # trailing newline

    for segs in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        shade(p, fill)
        if not segs:
            segs = [(" ", None, False)]
        for txt, color, italic in segs:
            r = p.add_run(txt)
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(size)
            r.italic = italic
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def read_text(path):
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


# ------------------------------------------------------------ running matlab
MATLAB_WRAPPER = r"""
try
    set(0,'DefaultFigureVisible','off');
    cd('{folder}');
    diary('{diary}');
    run('{script}');
    diary off;
    figs = findall(0,'type','figure');
    for k = 1:numel(figs)
        f = figs(k);
        if {fixplots}
            try
                set(f,'Color','w');
                axs = findall(f,'type','axes');
                for a = axs'
                    try, set(a,'Color','w','XColor','k','YColor','k','ZColor','k'); catch, end
                    try, set(a,'GridColor',[0.25 0.25 0.25],'MinorGridColor',[0.25 0.25 0.25],'GridAlpha',0.35); catch, end
                    try, set(a,'FontSize',max(a.FontSize,11)); catch, end
                    try, a.Title.Color='k'; a.Title.FontSize=14; a.Title.FontWeight='bold'; catch, end
                    try, a.XLabel.Color='k'; a.XLabel.FontSize=12; catch, end
                    try, a.YLabel.Color='k'; a.YLabel.FontSize=12; catch, end
                    try, a.ZLabel.Color='k'; a.ZLabel.FontSize=12; catch, end
                    for ln = findall(a,'type','line')'
                        try
                            c = ln.Color;
                            if mean(c) > 0.85, ln.Color = [0 0 0]; end
                        catch, end
                    end
                end
                try, set(findall(f,'type','legend'),'TextColor','k','Color','w','EdgeColor',[0.3 0.3 0.3]); catch, end
                try, set(findall(f,'type','colorbar'),'Color','k'); catch, end
                try, set(findall(f,'type','text'),'Color','k'); catch, end
            catch
            end
        end
        exportgraphics(f, fullfile('{outdir}', sprintf('fig_%02d.png', k)), 'Resolution', 150, 'BackgroundColor', 'white');
    end
catch ME
    diary off;
    fid = fopen('{diary}','a'); fprintf(fid, '\nERROR: %s\n', ME.message); fclose(fid);
end
exit;
"""

OCTAVE_WRAPPER = r"""
try
    cd('{folder}');
    diary('{diary}');
    run('{script}');
    diary off;
    figs = findall(0,'type','figure');
    for k = 1:numel(figs)
        f = figs(k);
        if {fixplots}
            try
                set(f,'color','w');
                set(findall(f,'type','axes'),'color','w','xcolor','k','ycolor','k','zcolor','k');
                set(findall(f,'type','text'),'color','k');
            catch
            end
        end
        print(f, fullfile('{outdir}', sprintf('fig_%02d.png', k)), '-dpng', '-r150');
    end
catch err
    diary off;
    fid = fopen('{diary}','a'); fprintf(fid, '\nERROR: %s\n', err.message); fclose(fid);
end
exit;
"""


def run_script(engine, mfile, timeout, fix_plots=True):
    """Run one .m file. Returns (text_output, [figure_paths], tmpdir)."""
    mfile = Path(mfile).resolve()   # absolute path, so cd works from any folder
    tmp = Path(tempfile.mkdtemp(prefix="mrep_"))
    diary = tmp / "out.txt"
    posix = lambda p: str(p).replace("\\", "/")
    wrapper = (MATLAB_WRAPPER if engine == "matlab" else OCTAVE_WRAPPER).format(
        folder=posix(mfile.parent), diary=posix(diary),
        script=mfile.name, outdir=posix(tmp),
        fixplots="1" if fix_plots else "0",
    )
    wfile = tmp / "wrapper.m"
    wfile.write_text(wrapper)
    if engine == "matlab":
        cmd = ["matlab", "-batch", f"run('{posix(wfile)}')"]
    else:
        cmd = ["octave", "--no-gui", "--quiet", str(wfile)]
    try:
        flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
        subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, creationflags=flags)
    except FileNotFoundError:
        raise RuntimeError(f"'{engine}' not found on PATH. Install it or choose not to run.")
    except subprocess.TimeoutExpired:
        diary.write_text("ERROR: timed out\n")
    text = read_text(diary) if diary.exists() else "(no output captured)"
    figs = sorted(tmp.glob("fig_*.png"))
    return text, figs, tmp


# -------------------------------------------------------------- repo scanning
def find_assignments(root):
    out = []
    for d in root.iterdir():
        if d.is_dir():
            m = re.search(r"assignment\W*(\d+)", d.name, re.I)
            if m:
                out.append((int(m.group(1)), d))
    return sorted(out)


def existing_outputs(folder, mfile):
    """Look for saved output next to the code (same stem) or generic output files."""
    stem = mfile.stem.lower()
    texts, imgs = [], []
    for f in sorted(folder.rglob("*"), key=natural_key):
        if not f.is_file():
            continue
        name = f.stem.lower()
        if f.suffix.lower() in IMG_EXT:
            imgs.append(f)
        elif f.suffix.lower() == ".txt" and (
            name.startswith(stem) or any(h in name for h in OUT_TXT_HINTS)
        ):
            texts.append(f)
    return texts, imgs


# ----------------------------------------------------------------------- main
def build(repo, out_path, engine, timeout=300, fix_plots=True, log=print, progress=lambda pct: None):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.9)

    assignments = find_assignments(repo)
    if not assignments:
        raise RuntimeError("No 'ASSIGNMENT N' folders found in " + str(repo))
    total = sum(max(1, len(list(f.rglob("*.m")))) for _, f in assignments)
    done = 0
    progress(0)

    for i, (num, folder) in enumerate(assignments):
        if i > 0:
            doc.add_page_break()
        log(f"== Assignment {num}")
        doc.add_heading(f"Assignment {num}", level=1)

        mfiles = sorted(folder.rglob("*.m"), key=natural_key)
        if not mfiles:
            doc.add_paragraph("(no .m files found)")
            done += 1
            progress(min(99, int(done / total * 100)))
            continue

        for mf in mfiles:
            log(f"   {mf.relative_to(repo)}")
            doc.add_heading(f"Code — {mf.name}", level=2)
            add_highlighted_code(doc, read_text(mf))

            doc.add_heading("Output", level=2)
            tmpdir = None
            if engine != "none":
                text, figs, tmpdir = run_script(engine, mf, timeout, fix_plots)
                texts = [text] if text.strip() else []
                imgs = figs
            else:
                tfiles, imgs = existing_outputs(mf.parent, mf)
                texts = [read_text(t) for t in tfiles]

            if not texts and not imgs:
                doc.add_paragraph("(no output found — run with --run matlab, or paste it here)")
            for t in texts:
                add_mono_block(doc, t.strip("\n"), size=9, fill="FFFFFF")
            for img in imgs:
                doc.add_paragraph()
                doc.add_picture(str(img), width=Inches(5.5))
            if tmpdir:
                # python-docx has already read the images by now
                shutil.rmtree(tmpdir, ignore_errors=True)
            doc.add_paragraph()
            done += 1
            progress(min(99, int(done / total * 100)))

    doc.save(out_path)
    progress(100)
    log(f"\nSaved: {out_path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--url", help="GitHub repo URL to clone")
    src.add_argument("--repo", help="Path to local clone")
    ap.add_argument("--out", default="MATLAB_Assignments.docx")
    ap.add_argument("--run", choices=["none", "matlab", "octave"], default="none",
                    help="Execute each .m file to capture output + figures")
    ap.add_argument("--no-fix-plots", action="store_true",
                    help="Keep figures exactly as MATLAB draws them (default: white background, dark readable text)")
    ap.add_argument("--timeout", type=int, default=300, help="Seconds per script when running")
    a = ap.parse_args()

    if a.url:
        repo = Path("_repo_tmp").resolve()
        if repo.exists():
            shutil.rmtree(repo)
        subprocess.run(["git", "clone", "--depth", "1", a.url, str(repo)], check=True)
    else:
        repo = Path(a.repo).resolve()
    try:
        build(repo, a.out, a.run, a.timeout, fix_plots=not a.no_fix_plots)
    except RuntimeError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()